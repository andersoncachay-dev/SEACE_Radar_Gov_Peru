"""Read side of Histórico Perú: filters, dashboard aggregates, paginated
lists, record detail and filtered Excel export."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date

from sqlalchemy import String, case, cast, distinct, func, or_, select
from sqlalchemy.orm import Session

from ...models import HistContract, HistContracting, HistExpiring, HistPacItem, HistRecurring
from .alerts_service import expiring_priority
from .derive import EXPIRING_BUCKETS, RECURRING_BUCKETS

MAX_PAGE_SIZE = 200
EXPORT_LIMIT = 100_000


@dataclass
class Filters:
    years: list[int] = field(default_factory=list)
    months: list[int] = field(default_factory=list)
    departments: list[str] = field(default_factory=list)
    categories: list[str] = field(default_factory=list)
    process_groups: list[str] = field(default_factory=list)
    government_levels: list[str] = field(default_factory=list)
    buckets: list[str] = field(default_factory=list)
    patterns: list[str] = field(default_factory=list)
    situations: list[str] = field(default_factory=list)
    q: str = ""
    relevance: str = "ok"
    real_date_only: bool = False
    hide_renewed: bool = False
    datasets: list[str] = field(default_factory=list)
    min_amount: float = 0
    # None = no favorites filter; otherwise the user's keys per favorite kind
    # (services/favorite_service.py), set by the router for "Solo favoritos".
    favorite_keys: dict[str, list[str]] | None = None


def _like(term: str) -> str:
    return f"%{term.strip()}%"


def _contracting_conditions(filters: Filters, model=HistContracting) -> list:
    conditions = []
    if filters.relevance in ("ok", "revisar"):
        conditions.append(model.relevance == filters.relevance)
    if filters.years:
        conditions.append(model.year.in_(filters.years))
    if filters.months:
        conditions.append(func.substr(model.month, 6, 2).in_([f"{m:02d}" for m in filters.months]))
    if filters.departments:
        conditions.append(model.department.in_(filters.departments))
    if filters.categories:
        conditions.append(model.category.in_(filters.categories))
    if filters.process_groups:
        conditions.append(model.process_group.in_(filters.process_groups))
    if filters.government_levels:
        conditions.append(model.government_level.in_(filters.government_levels))
    if filters.datasets:
        conditions.append(model.dataset.in_(filters.datasets))
    if filters.min_amount:
        conditions.append(model.amount >= filters.min_amount)
    if filters.favorite_keys is not None:
        conditions.append(model.source_key.in_(filters.favorite_keys.get("hist_contratacion", [])))
    if filters.q.strip():
        term = _like(filters.q)
        conditions.append(or_(
            model.entity_name.ilike(term), model.description.ilike(term), model.nomenclature.ilike(term),
            model.supplier_name.ilike(term), model.entity_ruc.ilike(term), model.supplier_ruc.ilike(term),
        ))
    return conditions


# --- filter options ------------------------------------------------------------------


def filter_options(db: Session) -> dict:
    def distinct_values(column, *where) -> list:
        return [value for value in db.scalars(select(distinct(column)).where(*where).order_by(column)).all() if value not in (None, "")]

    return {
        "years": distinct_values(HistContracting.year),
        "departments": distinct_values(HistContracting.department),
        "categories": distinct_values(HistContracting.category, HistContracting.relevance == "ok"),
        "process_groups": distinct_values(HistContracting.process_group),
        "government_levels": distinct_values(HistContracting.government_level),
        "expiring_buckets": list(EXPIRING_BUCKETS),
        "recurring_buckets": list(RECURRING_BUCKETS),
        "recurring_patterns": distinct_values(HistRecurring.pattern),
        "pac_situations": distinct_values(HistPacItem.situation),
        "pac_years": distinct_values(HistPacItem.year),
    }


# --- dashboard -----------------------------------------------------------------------------


def summary(db: Session, filters: Filters, today: date) -> dict:
    where = _contracting_conditions(filters)
    base = select(HistContracting).where(*where).subquery()

    def grouped(column, limit: int | None = None, order_by_amount: bool = True):
        query = select(column, func.coalesce(func.sum(base.c.amount), 0), func.count()).group_by(column)
        query = query.order_by(func.sum(base.c.amount).desc() if order_by_amount else column)
        if limit:
            query = query.limit(limit)
        return [{"key": key, "amount": round(float(amount or 0), 2), "count": count} for key, amount, count in db.execute(query).all()]

    totals = db.execute(
        select(func.coalesce(func.sum(base.c.amount), 0), func.count(), func.count(distinct(base.c.entity_ruc)), func.count(distinct(base.c.supplier_ruc)))
    ).one()
    top_entities = db.execute(
        select(base.c.entity_ruc, func.max(base.c.entity_name), func.max(base.c.department), func.sum(base.c.amount), func.count())
        .group_by(base.c.entity_ruc).order_by(func.sum(base.c.amount).desc()).limit(10)
    ).all()
    top_suppliers = db.execute(
        select(base.c.supplier_ruc, func.max(base.c.supplier_name), func.sum(base.c.amount), func.count(), func.count(distinct(base.c.entity_ruc)))
        .where(base.c.supplier_ruc != "").group_by(base.c.supplier_ruc).order_by(func.sum(base.c.amount).desc()).limit(10)
    ).all()
    return {
        "totals": {"amount": round(float(totals[0] or 0), 2), "count": totals[1], "entities": totals[2], "suppliers": totals[3]},
        "by_month": grouped(base.c.month, order_by_amount=False),
        "by_process_group": grouped(base.c.process_group),
        "by_category": grouped(base.c.category),
        "by_government_level": grouped(base.c.government_level),
        "by_department": grouped(base.c.department),
        "top_entities": [{"ruc": r, "name": n, "department": d, "amount": round(float(a or 0), 2), "count": c} for r, n, d, a, c in top_entities],
        "top_suppliers": [{"ruc": r, "name": n, "amount": round(float(a or 0), 2), "count": c, "entities": e} for r, n, a, c, e in top_suppliers],
        "expiring_buckets": bucket_counts(db, "vencimientos", filters),
        "recurring_buckets": bucket_counts(db, "recurrentes", filters),
        "runway": expiring_runway(db, filters, today),
    }


def bucket_counts(db: Session, kind: str, filters: Filters) -> list[dict]:
    if kind == "vencimientos":
        query = (
            select(HistExpiring.bucket, func.count(), func.coalesce(func.sum(HistContracting.amount), 0))
            .join(HistContracting, HistContracting.id == HistExpiring.contracting_id)
            .where(*_expiring_conditions(_without(filters, "years", "months", "buckets")))
            .group_by(HistExpiring.bucket)
        )
        order = EXPIRING_BUCKETS
    else:
        query = select(HistRecurring.bucket, func.count(), func.coalesce(func.sum(HistRecurring.amount_12m_supplier), 0)).where(
            *_recurring_conditions(_without(filters, "buckets"))
        ).group_by(HistRecurring.bucket)
        order = RECURRING_BUCKETS
    found = {bucket: (count, amount) for bucket, count, amount in db.execute(query).all()}
    return [{"key": bucket, "count": found.get(bucket, (0, 0))[0], "amount": round(float(found.get(bucket, (0, 0))[1]), 2)} for bucket in order]


def _without(filters: Filters, *names: str) -> Filters:
    copy = Filters(**{k: (list(v) if isinstance(v, list) else v) for k, v in filters.__dict__.items()})
    for name in names:
        setattr(copy, name, [])
    return copy


def expiring_runway(db: Session, filters: Filters, today: date, limit: int = 600) -> list[dict]:
    """Contracts ending in the next 12 months, for the "pista de vencimientos"."""
    rows = db.execute(
        select(HistExpiring, HistContracting)
        .join(HistContracting, HistContracting.id == HistExpiring.contracting_id)
        .where(HistExpiring.days_to_end >= 0, HistExpiring.days_to_end <= 365, *_expiring_conditions(_without(filters, "years", "months")))
        .order_by(HistContracting.amount.desc())
        .limit(limit)
    ).all()
    return [{
        "id": expiring.id,
        "end_date": expiring.end_date.isoformat(),
        "days": expiring.days_to_end,
        "amount": award.amount,
        "entity_name": award.entity_name,
        "department": award.department,
        "is_real_date": expiring.is_real_date,
        "possibly_renewed": expiring.possibly_renewed,
        "priority": expiring_priority(expiring.days_to_end, award.amount) or "P4",
    } for expiring, award in rows]


# --- lists ---------------------------------------------------------------------------------


def _page(page: int, page_size: int) -> tuple[int, int]:
    size = min(max(page_size, 1), MAX_PAGE_SIZE)
    return max(page, 1), size


def _contracting_dict(row: HistContracting) -> dict:
    return {
        "id": row.id, "dataset": row.dataset, "year": row.year, "month": row.month,
        "date": row.event_date.isoformat() if row.event_date else None,
        "process_group": row.process_group, "procedure": row.procedure, "nomenclature": row.nomenclature,
        "entity_ruc": row.entity_ruc, "entity_name": row.entity_name, "department": row.department,
        "description": row.description, "category": row.category, "amount": row.amount,
        "reference_amount": row.reference_amount, "supplier_ruc": row.supplier_ruc, "supplier_name": row.supplier_name,
        "status": row.status, "direct_cause": row.direct_cause, "object_type": row.object_type,
        "government_level": row.government_level, "relevance": row.relevance,
        "fav_key": row.source_key,
    }


CONTRACTING_SORTS = {
    "date_desc": (HistContracting.event_date.desc(), HistContracting.id.desc()),
    "date_asc": (HistContracting.event_date.asc(), HistContracting.id.asc()),
    "amount_desc": (HistContracting.amount.desc(), HistContracting.id.desc()),
}


def list_contractings(db: Session, filters: Filters, page: int = 1, page_size: int = 50, sort: str = "date_desc") -> dict:
    page, size = _page(page, page_size)
    where = _contracting_conditions(filters)
    total, amount = db.execute(select(func.count(), func.coalesce(func.sum(HistContracting.amount), 0)).where(*where)).one()
    rows = db.scalars(
        select(HistContracting).where(*where).order_by(*CONTRACTING_SORTS.get(sort, CONTRACTING_SORTS["date_desc"])).offset((page - 1) * size).limit(size)
    ).all()
    return {"items": [_contracting_dict(row) for row in rows], "total": total, "amount": round(float(amount), 2), "page": page, "page_size": size}


def _expiring_conditions(filters: Filters) -> list:
    conditions = _contracting_conditions(filters)
    if filters.buckets:
        conditions.append(HistExpiring.bucket.in_(filters.buckets))
    if filters.real_date_only:
        conditions.append(HistExpiring.is_real_date.is_(True))
    if filters.hide_renewed:
        conditions.append(HistExpiring.possibly_renewed.is_(False))
    return conditions


def _expiring_dict(expiring: HistExpiring, award: HistContracting, contract: HistContract | None) -> dict:
    return {
        **_contracting_dict(award),
        "id": expiring.id,
        "contracting_id": award.id,
        "start_date": expiring.start_date.isoformat() if expiring.start_date else None,
        "end_date": expiring.end_date.isoformat(),
        "date_source": expiring.date_source,
        "is_real_date": expiring.is_real_date,
        "days_to_end": expiring.days_to_end,
        "bucket": expiring.bucket,
        "monthly_amount": expiring.monthly_amount,
        "extras_amount": expiring.extras_amount,
        "possibly_renewed": expiring.possibly_renewed,
        "priority": expiring_priority(expiring.days_to_end, award.amount),
        "contract_number": contract.contract_number if contract else "",
        "contract_url": contract.url if contract else "",
    }


EXPIRING_SORTS = {
    "days_asc": (HistExpiring.days_to_end.asc(),),
    "amount_desc": (HistContracting.amount.desc(),),
    "days_desc": (HistExpiring.days_to_end.desc(),),
}


def _expiring_query(filters: Filters):
    return (
        select(HistExpiring, HistContracting, HistContract)
        .join(HistContracting, HistContracting.id == HistExpiring.contracting_id)
        .outerjoin(HistContract, HistContract.id == HistExpiring.contract_id)
        .where(*_expiring_conditions(filters))
    )


def list_expiring(db: Session, filters: Filters, page: int = 1, page_size: int = 50, sort: str = "days_asc") -> dict:
    page, size = _page(page, page_size)
    where = _expiring_conditions(filters)
    total, amount = db.execute(
        select(func.count(), func.coalesce(func.sum(HistContracting.amount), 0)).select_from(HistExpiring)
        .join(HistContracting, HistContracting.id == HistExpiring.contracting_id).where(*where)
    ).one()
    order = EXPIRING_SORTS.get(sort, EXPIRING_SORTS["days_asc"])
    rows = db.execute(_expiring_query(filters).order_by(*order, HistExpiring.id).offset((page - 1) * size).limit(size)).all()
    return {"items": [_expiring_dict(*row) for row in rows], "total": total, "amount": round(float(amount), 2), "page": page, "page_size": size}


def recurring_fav_key():
    """SQL form of the hist_recurrente favorite key (same as its alert ref_key)."""
    return HistRecurring.entity_ruc + "|" + HistRecurring.supplier_ruc


def pac_fav_key():
    """SQL form of the hist_pac favorite key (same as its alert ref_key)."""
    return cast(HistPacItem.year, String) + "|" + HistPacItem.entity_code + "|" + HistPacItem.reference_number + "|" + HistPacItem.item_number


def _recurring_conditions(filters: Filters) -> list:
    conditions = []
    if filters.departments:
        conditions.append(HistRecurring.department.in_(filters.departments))
    if filters.categories:
        conditions.append(HistRecurring.category.in_(filters.categories))
    if filters.government_levels:
        conditions.append(HistRecurring.government_level.in_(filters.government_levels))
    if filters.buckets:
        conditions.append(HistRecurring.bucket.in_(filters.buckets))
    if filters.patterns:
        conditions.append(HistRecurring.pattern.in_(filters.patterns))
    if filters.min_amount:
        conditions.append(HistRecurring.amount_12m_supplier >= filters.min_amount)
    if filters.favorite_keys is not None:
        conditions.append(recurring_fav_key().in_(filters.favorite_keys.get("hist_recurrente", [])))
    if filters.q.strip():
        term = _like(filters.q)
        conditions.append(or_(HistRecurring.entity_name.ilike(term), HistRecurring.supplier_name.ilike(term), HistRecurring.last_order_description.ilike(term)))
    return conditions


def _recurring_dict(row: HistRecurring) -> dict:
    return {
        "id": row.id, "entity_ruc": row.entity_ruc, "entity_name": row.entity_name, "department": row.department,
        "government_level": row.government_level, "supplier_ruc": row.supplier_ruc, "supplier_name": row.supplier_name,
        "category": row.category, "orders_12m": row.orders_12m, "months_12m": row.months_12m,
        "amount_12m_supplier": row.amount_12m_supplier, "amount_12m_entity": row.amount_12m_entity, "suppliers_12m": row.suppliers_12m,
        "last_order_date": row.last_order_date.isoformat() if row.last_order_date else None,
        "last_order_amount": row.last_order_amount, "last_order_description": row.last_order_description,
        "pattern": row.pattern, "next_purchase_date": row.next_purchase_date.isoformat() if row.next_purchase_date else None,
        "days_to_next": row.days_to_next, "bucket": row.bucket,
        "fav_key": f"{row.entity_ruc}|{row.supplier_ruc}",
    }


RECURRING_SORTS = {
    "days_asc": (HistRecurring.days_to_next.asc(),),
    "amount_desc": (HistRecurring.amount_12m_supplier.desc(),),
}


def list_recurring(db: Session, filters: Filters, page: int = 1, page_size: int = 50, sort: str = "days_asc") -> dict:
    page, size = _page(page, page_size)
    where = _recurring_conditions(filters)
    total, amount = db.execute(select(func.count(), func.coalesce(func.sum(HistRecurring.amount_12m_supplier), 0)).where(*where)).one()
    rows = db.scalars(
        select(HistRecurring).where(*where).order_by(*RECURRING_SORTS.get(sort, RECURRING_SORTS["days_asc"]), HistRecurring.id).offset((page - 1) * size).limit(size)
    ).all()
    return {"items": [_recurring_dict(row) for row in rows], "total": total, "amount": round(float(amount), 2), "page": page, "page_size": size}


def _pac_conditions(filters: Filters) -> list:
    conditions = [HistPacItem.relevance == "ok"]
    if filters.years:
        conditions.append(HistPacItem.year.in_(filters.years))
    if filters.months:
        conditions.append(HistPacItem.planned_month.in_(filters.months))
    if filters.departments:
        conditions.append(HistPacItem.department.in_(filters.departments))
    if filters.categories:
        conditions.append(HistPacItem.category.in_(filters.categories))
    if filters.process_groups:
        conditions.append(HistPacItem.process_group.in_(filters.process_groups))
    if filters.government_levels:
        conditions.append(HistPacItem.government_level.in_(filters.government_levels))
    if filters.situations:
        conditions.append(HistPacItem.situation.in_(filters.situations))
    if filters.favorite_keys is not None:
        conditions.append(pac_fav_key().in_(filters.favorite_keys.get("hist_pac", [])))
    if filters.q.strip():
        term = _like(filters.q)
        conditions.append(or_(HistPacItem.entity_name.ilike(term), HistPacItem.item_description.ilike(term), HistPacItem.process_description.ilike(term)))
    return conditions


def pac_summary(db: Session, filters: Filters) -> dict:
    """PAC dashboard: totals, items per planned month, process type,
    department and situation, all under the same filters as the list."""
    where = _pac_conditions(filters)

    def grouped(column, order_by_count: bool = True):
        query = select(column, func.count()).where(*where).group_by(column)
        query = query.order_by(func.count().desc() if order_by_count else column)
        return [{"key": key, "count": count} for key, count in db.execute(query).all()]

    totals = db.execute(
        select(func.count(), func.count(distinct(HistPacItem.entity_ruc)), func.coalesce(func.sum(case((HistPacItem.already_awarded.is_(True), 1), else_=0)), 0)).where(*where)
    ).one()
    open_items = db.scalar(select(func.count()).select_from(HistPacItem).where(*where, HistPacItem.situation.like("Por convocar%")))
    return {
        "totals": {"items": totals[0], "entities": totals[1], "already_awarded": int(totals[2] or 0), "to_call": open_items},
        "by_month": grouped(HistPacItem.planned_month, order_by_count=False),
        "by_process_group": grouped(HistPacItem.process_group),
        "by_department": grouped(HistPacItem.department),
        "by_situation": grouped(HistPacItem.situation),
        "by_category": grouped(HistPacItem.category),
    }


def _pac_dict(row: HistPacItem) -> dict:
    return {
        "id": row.id, "year": row.year, "entity_ruc": row.entity_ruc, "entity_name": row.entity_name, "department": row.department,
        "government_level": row.government_level, "version": row.version, "reference_number": row.reference_number,
        "process_description": row.process_description, "process_type": row.process_type, "process_group": row.process_group,
        "object_type": row.object_type, "item_number": row.item_number, "item_description": row.item_description,
        "category": row.category, "quantity": row.quantity, "unit": row.unit, "item_department": row.item_department,
        "province": row.province, "district": row.district, "planned_month": row.planned_month,
        "published_at": row.published_at.isoformat() if row.published_at else None,
        "situation": row.situation, "already_awarded": row.already_awarded,
        "fav_key": f"{row.year}|{row.entity_code}|{row.reference_number}|{row.item_number}",
    }


def list_pac(db: Session, filters: Filters, page: int = 1, page_size: int = 50, sort: str = "month_asc") -> dict:
    page, size = _page(page, page_size)
    where = _pac_conditions(filters)
    total = db.scalar(select(func.count()).select_from(HistPacItem).where(*where))
    order = (HistPacItem.planned_month.desc(),) if sort == "month_desc" else (HistPacItem.planned_month.asc(),)
    rows = db.scalars(select(HistPacItem).where(*where).order_by(*order, HistPacItem.id).offset((page - 1) * size).limit(size)).all()
    return {"items": [_pac_dict(row) for row in rows], "total": total, "amount": 0, "page": page, "page_size": size}


# --- detail --------------------------------------------------------------------------------


def _entity_history(db: Session, entity_ruc: str) -> list[dict]:
    rows = db.execute(
        select(HistContracting.year, func.sum(HistContracting.amount), func.count())
        .where(HistContracting.entity_ruc == entity_ruc, HistContracting.relevance == "ok")
        .group_by(HistContracting.year).order_by(HistContracting.year)
    ).all()
    return [{"year": year, "amount": round(float(amount or 0), 2), "count": count} for year, amount, count in rows]


def _entity_suppliers(db: Session, entity_ruc: str) -> list[dict]:
    rows = db.execute(
        select(HistContracting.supplier_ruc, func.max(HistContracting.supplier_name), func.sum(HistContracting.amount), func.count(), func.max(HistContracting.event_date))
        .where(HistContracting.entity_ruc == entity_ruc, HistContracting.relevance == "ok", HistContracting.supplier_ruc != "")
        .group_by(HistContracting.supplier_ruc).order_by(func.sum(HistContracting.amount).desc()).limit(8)
    ).all()
    return [{"ruc": r, "name": n, "amount": round(float(a or 0), 2), "count": c, "last_date": d.isoformat() if d else None} for r, n, a, c, d in rows]


def _related_awards(db: Session, entity_ruc: str, exclude_id: int) -> list[dict]:
    rows = db.scalars(
        select(HistContracting).where(
            HistContracting.entity_ruc == entity_ruc, HistContracting.dataset == "adjudicacion",
            HistContracting.relevance == "ok", HistContracting.id != exclude_id,
        ).order_by(HistContracting.event_date.desc()).limit(8)
    ).all()
    return [_contracting_dict(row) for row in rows]


def detail(db: Session, kind: str, record_id: int) -> dict | None:
    if kind == "contrataciones":
        row = db.get(HistContracting, record_id)
        if not row:
            return None
        expiring = db.scalar(select(HistExpiring).where(HistExpiring.contracting_id == row.id))
        record = _contracting_dict(row)
        if expiring:
            contract = db.get(HistContract, expiring.contract_id) if expiring.contract_id else None
            record = _expiring_dict(expiring, row, contract) | {"id": row.id, "expiring_id": expiring.id}
        entity_ruc, exclude = row.entity_ruc, row.id
    elif kind == "vencimientos":
        found = db.execute(_expiring_query(Filters(relevance="")).where(HistExpiring.id == record_id)).first()
        if not found:
            return None
        record = _expiring_dict(*found)
        entity_ruc, exclude = found[1].entity_ruc, found[1].id
    elif kind == "recurrentes":
        row = db.get(HistRecurring, record_id)
        if not row:
            return None
        record = _recurring_dict(row)
        orders = db.scalars(
            select(HistContracting).where(HistContracting.dataset == "orden", HistContracting.entity_ruc == row.entity_ruc, HistContracting.supplier_ruc == row.supplier_ruc)
            .order_by(HistContracting.event_date.desc()).limit(24)
        ).all()
        record["orders"] = [_contracting_dict(order) for order in orders]
        entity_ruc, exclude = row.entity_ruc, 0
    elif kind == "pac":
        row = db.get(HistPacItem, record_id)
        if not row:
            return None
        record = _pac_dict(row)
        entity_ruc, exclude = row.entity_ruc, 0
    else:
        return None
    return {
        "kind": kind,
        "record": record,
        "entity_history": _entity_history(db, entity_ruc),
        "entity_suppliers": _entity_suppliers(db, entity_ruc),
        "related_awards": _related_awards(db, entity_ruc, exclude),
    }


# --- export --------------------------------------------------------------------------------

EXPORT_COLUMNS: dict[str, list[tuple[str, str]]] = {
    "contrataciones": [
        ("year", "Año"), ("month", "Mes"), ("date", "Fecha"), ("process_group", "Tipo de proceso"), ("procedure", "Procedimiento SEACE"),
        ("nomenclature", "Nomenclatura / Orden"), ("entity_ruc", "RUC Entidad"), ("entity_name", "Entidad"), ("department", "Departamento"),
        ("description", "Descripción"), ("category", "Categoría"), ("amount", "Monto S/"), ("reference_amount", "Monto referencial S/"),
        ("supplier_ruc", "RUC Proveedor"), ("supplier_name", "Proveedor"), ("status", "Estado"), ("direct_cause", "Causal (Contratación Directa)"),
        ("object_type", "Objeto"), ("government_level", "Nivel de gobierno"),
    ],
    "vencimientos": [
        ("bucket", "Estado"), ("days_to_end", "Días para vencer"), ("end_date", "Fin de vigencia"), ("date_source", "Fuente de la fecha"),
        ("entity_name", "Entidad"), ("department", "Departamento"), ("process_group", "Tipo de proceso"), ("nomenclature", "Nomenclatura"),
        ("description", "Descripción"), ("category", "Categoría"), ("supplier_name", "Proveedor actual"), ("supplier_ruc", "RUC proveedor"),
        ("amount", "Monto adjudicado S/"), ("monthly_amount", "Monto mensual aprox. S/"), ("extras_amount", "Adicionales/prórrogas S/"),
        ("start_date", "Inicio"), ("date", "Buena pro"), ("contract_number", "Nº contrato"), ("possibly_renewed", "Posible renovado"),
        ("direct_cause", "Causal (CD)"), ("entity_ruc", "RUC entidad"), ("contract_url", "URL contrato"), ("government_level", "Nivel de gobierno"),
    ],
    "recurrentes": [
        ("bucket", "Ventana próxima compra"), ("next_purchase_date", "Próxima compra estimada"), ("days_to_next", "Días"), ("pattern", "Patrón"),
        ("entity_name", "Entidad"), ("department", "Departamento"), ("supplier_name", "Proveedor actual"), ("orders_12m", "Órdenes 12m"),
        ("months_12m", "Meses con orden 12m"), ("amount_12m_supplier", "Monto 12m S/ (este proveedor)"), ("amount_12m_entity", "Monto 12m S/ (entidad, todos)"),
        ("suppliers_12m", "Proveedores de la entidad 12m"), ("last_order_date", "Última orden"), ("last_order_amount", "Monto última orden S/"),
        ("last_order_description", "Descripción última orden"), ("category", "Categoría"), ("entity_ruc", "RUC entidad"), ("supplier_ruc", "RUC proveedor"),
        ("government_level", "Nivel de gobierno"),
    ],
    "pac": [
        ("situation", "Situación"), ("planned_month", "Mes previsto"), ("entity_name", "Entidad"), ("department", "Departamento"),
        ("process_type", "Tipo de proceso"), ("object_type", "Objeto"), ("process_description", "Descripción del proceso"), ("item_description", "Ítem"),
        ("quantity", "Cantidad"), ("unit", "Unidad"), ("item_department", "Lugar del ítem (depto)"), ("province", "Provincia"), ("district", "Distrito"),
        ("already_awarded", "Ya adjudicó en el año"), ("reference_number", "Nº referencia PAC"), ("version", "Versión PAC"), ("entity_ruc", "RUC entidad"),
        ("government_level", "Nivel de gobierno"),
    ],
}

_LISTERS = {"contrataciones": list_contractings, "vencimientos": list_expiring, "recurrentes": list_recurring, "pac": list_pac}


def export_xlsx(db: Session, kind: str, filters: Filters, sort: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    lister = _LISTERS[kind]
    columns = EXPORT_COLUMNS[kind]
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=kind.capitalize()[:31])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0A2449")
    from openpyxl.cell import WriteOnlyCell

    header = []
    for _, label in columns:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font, cell.fill = header_font, header_fill
        header.append(cell)
    sheet.append(header)
    page, written = 1, 0
    while written < EXPORT_LIMIT:
        kwargs = {"page": page, "page_size": MAX_PAGE_SIZE}
        if sort:
            kwargs["sort"] = sort
        batch = lister(db, filters, **kwargs)["items"]
        if not batch:
            break
        for item in batch:
            sheet.append([_excel_value(item.get(key)) for key, _ in columns])
        written += len(batch)
        page += 1
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _excel_value(value):
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, str) and len(value) == 10 and value[4] == "-" and value[7] == "-":
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value
